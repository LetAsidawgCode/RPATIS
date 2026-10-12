import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOCX_DIR = r"c:\Users\User\Desktop\Assesors Office\Tools\DOCX"
INDIV_DIR = os.path.join(DOCX_DIR, "Individual_Files")
ETC_DIR = r"c:\Users\User\Desktop\Assesors Office\Tools\Etc"
os.makedirs(DOCX_DIR, exist_ok=True)
os.makedirs(INDIV_DIR, exist_ok=True)

COLOR_MAROON = RGBColor(122, 17, 19)      # #7A1113
COLOR_DARK = RGBColor(34, 34, 34)          # #222222
COLOR_MUTED = RGBColor(85, 85, 85)        # #555555
HEX_MAROON = "7A1113"
HEX_LIGHT_BG = "F8F9FA"
HEX_LIGHT_RED = "FDF7F7"
HEX_BORDER = "CED4DA"
HEX_MAROON_BORDER = "7A1113"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D5DD", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_cell_text(cell, text, bold=False, italic=False, size=8, color=COLOR_DARK, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.name = "Arial"
    return p

def add_header(doc, doc_title, doc_sub="", step_tag=""):
    p_u1 = doc.add_paragraph()
    p_u1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_u1.paragraph_format.space_before = Pt(0)
    p_u1.paragraph_format.space_after = Pt(1)
    r1 = p_u1.add_run("UNIVERSITY OF NUEVA CACERES")
    r1.bold = True
    r1.font.size = Pt(11)
    r1.font.color.rgb = COLOR_MAROON
    r1.font.name = "Arial"

    p_u2 = doc.add_paragraph()
    p_u2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_u2.paragraph_format.space_before = Pt(0)
    p_u2.paragraph_format.space_after = Pt(1)
    r2 = p_u2.add_run("School of Computer and Information Sciences")
    r2.font.size = Pt(9)
    r2.font.color.rgb = COLOR_DARK
    r2.font.name = "Arial"

    p_u3 = doc.add_paragraph()
    p_u3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_u3.paragraph_format.space_before = Pt(0)
    p_u3.paragraph_format.space_after = Pt(4)
    r3 = p_u3.add_run("Bachelor of Science in Information Technology • Naga City, Camarines Sur")
    r3.font.size = Pt(8)
    r3.font.color.rgb = COLOR_MUTED
    r3.font.name = "Arial"

    h1 = doc.add_heading(level=1)
    h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(2)
    h1.paragraph_format.space_after = Pt(2)
    full_title = f"{step_tag}: {doc_title}" if step_tag else doc_title
    hrun = h1.add_run(full_title.upper())
    hrun.bold = True
    hrun.font.size = Pt(12)
    hrun.font.color.rgb = COLOR_MAROON
    hrun.font.name = "Arial"

    if doc_sub:
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(6)
        r_sub = p_sub.add_run(doc_sub)
        r_sub.italic = True
        r_sub.font.size = Pt(8.5)
        r_sub.font.color.rgb = COLOR_MUTED
        r_sub.font.name = "Arial"

def add_callout(doc, title, lines):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    set_cell_background(cell, HEX_LIGHT_BG)
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    if title:
        r_t = p.add_run(title)
        r_t.bold = True
        r_t.font.size = Pt(8.5)
        r_t.font.color.rgb = COLOR_MAROON
        r_t.font.name = "Arial"

    for idx, l in enumerate(lines):
        p_l = cell.add_paragraph() if (title or idx > 0) else p
        p_l.paragraph_format.space_before = Pt(0)
        p_l.paragraph_format.space_after = Pt(1.5)
        r_l = p_l.add_run(l)
        r_l.font.size = Pt(8)
        r_l.font.color.rgb = COLOR_DARK
        r_l.font.name = "Arial"

    doc.add_paragraph().paragraph_format.space_after = Pt(3)

def add_banner(doc, text):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    set_cell_background(cell, HEX_MAROON)
    set_cell_margins(cell, top=60, bottom=60, left=120, right=120)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(255, 255, 255)
    r.font.name = "Arial"
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def build_tool_01_content(doc):
    add_header(doc, "Expert Validation Sheet for Research Instruments",
               "Step 1: Content Validity, Construct Adequacy, and Methodological Rigor Evaluation Protocol",
               "Step 1")
    add_callout(doc, "RESEARCH TIMELINE STAGE: STEP 1 (PRE-DEPLOYMENT VALIDATION)", [
        "When to Use: Submit this form to expert validators before field deployment to formally evaluate and certify the research tools.",
        "Study Title: Needs Assessment for a Real Property Assessment Tracking and Inter-Office Synchronization System: Resolving Disconnected Records and Workflow Delays in Municipal Assessor and Treasurer Offices in Pasacao",
        "Student Researchers: Rizalina A. Ortega | John Kenneth P. Calubad | Christian Louise O. Asido",
        "Instruments Evaluated: Step 2 Consent Form, Step 3 Interview Guide, Step 4 Time-Motion Log, Step 5 ISO 25010 Survey."
    ])

    # Validator Details Grid
    t_v = doc.add_table(rows=3, cols=2)
    t_v.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_v)
    format_cell_text(t_v.cell(0, 0), "Validator Name: ___________________________", bold=True)
    format_cell_text(t_v.cell(0, 1), "Date of Evaluation: ___________________________", bold=True)
    format_cell_text(t_v.cell(1, 0), "Position / Designation: ___________________________")
    format_cell_text(t_v.cell(1, 1), "Institutional Affiliation: ___________________________")
    format_cell_text(t_v.cell(2, 0), "Area of Expertise:\n[  ] Information Technology / Software Engineering\n[  ] Public Administration / LGU\n[  ] Academic Research / Statistics\n[  ] Other: ___________________________")
    format_cell_text(t_v.cell(2, 1), "Highest Degree:\n[  ] Master's Degree (MS / MIT / MPA)\n[  ] Doctorate Degree (PhD / DIT / DPA)\n[  ] Professional License (REA / REB / CPA)\nYears in Field: ______ years")
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    add_callout(doc, "RATING SCALE & BENCHMARK", [
        "5 = Strongly Agree (Exceptionally satisfied)  |  4 = Agree (Adequately satisfied)  |  3 = Neutral (Revisions suggested)",
        "2 = Disagree (Substantial restructuring needed)  |  1 = Strongly Disagree (Invalid/Inadequate)"
    ])

    add_banner(doc, "Instrument Validation Matrix (Criteria Across Tools)")
    headers = ["#", "Validation Evaluation Criterion", "Step 02 (Consent)", "Step 03 (Interview)", "Step 04 (Time-Motion)", "Step 05 (ISO 25010)"]
    criteria = [
        ("1", "Construct & Objective Alignment: Items directly measure the operational variables and objectives specified in Chapter 1."),
        ("2", "Clarity and Unambiguity: Questions, rating indicators, and instructions are phrased clearly and free of confusing wording."),
        ("3", "Language Appropriateness: Vocabulary and phrasing match the professional profile of municipal and provincial personnel."),
        ("4", "Content Comprehensiveness: The tools comprehensively cover administrative workflows, timing metrics, and software quality."),
        ("5", "Objectivity and Non-Bias: Items are neutral, balanced, and free from leading questions or presupposed outcomes."),
        ("6", "Measurement Scale Suitability: Response options, timestamp entries, and Likert scales are operationally and statistically sound."),
        ("7", "Instrument Economy: Length and completion duration are realistic and respectful of municipal operating hours."),
        ("8", "Ethical and Statutory Soundness: Conforms strictly to the Data Privacy Act (RA 10173) and UNC ethics guidelines."),
    ]
    t_mat = doc.add_table(rows=len(criteria) + 2, cols=6)
    t_mat.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_mat)
    
    # Headers
    for c_idx, h_text in enumerate(headers):
        cell = t_mat.cell(0, c_idx)
        set_cell_background(cell, HEX_MAROON)
        format_cell_text(cell, h_text, bold=True, size=7.5, color=RGBColor(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)

    # Rows
    for r_idx, (num, text) in enumerate(criteria, start=1):
        format_cell_text(t_mat.cell(r_idx, 0), num, size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell_text(t_mat.cell(r_idx, 1), text, size=7.5)
        for c in range(2, 6):
            format_cell_text(t_mat.cell(r_idx, c), "[   ]", size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Average row
    avg_row = len(criteria) + 1
    format_cell_text(t_mat.cell(avg_row, 0), "∑", bold=True, size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    format_cell_text(t_mat.cell(avg_row, 1), "COMPOSITE MEAN RATING", bold=True, size=7.5)
    for c in range(2, 6):
        format_cell_text(t_mat.cell(avg_row, c), "______", bold=True, size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    add_callout(doc, "VALIDATION VERDICT & SPECIFIC RECOMMENDATIONS", [
        "Step 02 (Consent Form): _____________________________________________________________________________",
        "Step 03 (Interview Guide): __________________________________________________________________________",
        "Step 04 (Time-Motion Log): __________________________________________________________________________",
        "Step 05 (ISO 25010 Survey): _________________________________________________________________________",
        "OVERALL VERDICT:  [  ] Approved for Field Use    [  ] Approved with Minor Revisions    [  ] Needs Major Revisions",
        "Validator Signature: ____________________________________        Date: ________________________",
        "Printed Name: __________________________________________        Official Stamp / Seal: ________________"
    ])

def build_tool_02_content(doc):
    add_header(doc, "Participant Consent and Data Privacy Agreement Form",
               "Step 2: Republic Act No. 10173 Compliance & UNC Institutional Ethics Protocol",
               "Step 2")
    add_callout(doc, "RESEARCH TIMELINE STAGE: STEP 2 (PARTICIPANT ETHICAL CLEARANCE)", [
        "When to Use: Must be signed by municipal personnel and evaluators before any interview, timing observation, or system testing.",
        "Project Title: Needs Assessment for a Real Property Assessment Tracking and Inter-Office Synchronization System: Resolving Disconnected Records and Workflow Delays in Municipal Assessor and Treasurer Offices in Pasacao",
        "Student Researchers: Rizalina A. Ortega | John Kenneth P. Calubad | Christian Louise O. Asido",
        "Academic Department: School of Computer and Information Sciences, University of Nueva Caceres"
    ])

    sections = [
        ("1. Purpose and Nature of the Research",
         "The research team is conducting a study to analyze operational workflows, administrative disconnects, and document routing latency between the Municipal Assessor's Office, the Municipal Treasurer's Office, and the Provincial Assessor's Office in Pili. The ultimate objective is to design, develop, and evaluate a web-based Real Property Assessment Tracking and Inter-Office Synchronization System that eliminates manual cross-checking and streamlines document routing."),
        
        ("2. Participant Roles and Required Activities",
         "As a designated municipal or provincial official, administrative staff member, or technical expert, you may be invited to participate in one or more of the following evaluation activities:\n"
         "• Stakeholder Needs Assessment Interview: Answering semi-structured questions regarding daily office workflows, document verification, legacy archives, and inter-office bottlenecks (approx. 20–30 minutes).\n"
         "• Time-Motion Workflow Observation: Allowing the research team to observe and record chronological processing durations for routine assessment and reconciliation tasks (approx. 15–30 minutes).\n"
         "• Software Pilot Testing & ISO/IEC 25010 Evaluation: Testing the prototype web application and completing an evaluation survey assessing functional suitability, efficiency, usability, security, and reliability (approx. 20–30 minutes)."),
        
        ("3. Voluntary Participation and Right of Withdrawal",
         "Your participation in this research study is entirely voluntary. You reserve the absolute right to refuse to answer any individual question, decline to participate in specific procedures, or terminate your participation at any time without penalty, prejudice, or impact on your official employment status."),
        
        ("4. Strict Confidentiality and Data Privacy Protections (RA 10173)",
         "In strict compliance with Republic Act No. 10173 (Data Privacy Act of 2012), the research team enforces rigorous data protection safeguards:\n"
         "• Strict Anonymity: Published findings, academic manuscripts, and statistical tables will strictly report aggregate metrics and anonymized designations. Personal names will never appear.\n"
         "• Zero Exposure of Citizen Records: During development, testing, and simulation trials, only masked or synthetic dummy property records will be utilized. No actual private taxpayer financial data or confidential titles will be recorded.\n"
         "• Data Security & Retention: All qualitative notes, observation sheets, and survey data will be securely stored in encrypted academic repositories accessible solely to the research proponents and their academic adviser, and will be destroyed following final project completion.\n"
         "• Audio Recording Consent: Audio recording during interviews is strictly optional and requires explicit affirmative consent."),
        
        ("5. Assessment of Risks and Anticipated Benefits",
         "There are no foreseeable physical, legal, or psychological risks associated with this study beyond routine professional office interactions. Your expert contributions will directly shape a software intervention designed to eliminate tedious manual paper handoffs, shorten transaction delays, and improve frontline public services for local taxpayers."),
        
        ("6. Research Inquiries and Institutional Contact",
         "If you have inquiries, concerns, or require clarification regarding this study or your rights under the Data Privacy Act, you may contact the student researchers or the Capstone Committee via the School of Computer and Information Sciences, University of Nueva Caceres, Naga City."),
    ]

    for title, body in sections:
        p_t = doc.add_paragraph()
        p_t.paragraph_format.space_before = Pt(2)
        p_t.paragraph_format.space_after = Pt(1)
        r_t = p_t.add_run(title)
        r_t.bold = True
        r_t.font.size = Pt(8.5)
        r_t.font.color.rgb = COLOR_DARK
        r_t.font.name = "Arial"

        p_b = doc.add_paragraph()
        p_b.paragraph_format.space_before = Pt(0)
        p_b.paragraph_format.space_after = Pt(3)
        r_b = p_b.add_run(body)
        r_b.font.size = Pt(8)
        r_b.font.color.rgb = COLOR_DARK
        r_b.font.name = "Arial"

    add_banner(doc, "Affirmative Participant Consent and Authorization")
    t_c = doc.add_table(rows=3, cols=2)
    t_c.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_c)
    
    # Span top row
    t_c.cell(0, 0).merge(t_c.cell(0, 1))
    format_cell_text(t_c.cell(0, 0), "I have read and understood the information provided in this Informed Consent and Data Privacy Agreement Form. I have been afforded sufficient opportunity to ask questions, which were answered to my satisfaction. I voluntarily agree to participate in this academic research study.\n\nAudio Recording Authorization:\n[  ] YES, I consent to audio recording for note-taking purposes only.\n[  ] NO, I do not consent to audio recording; written notes only.")
    format_cell_text(t_c.cell(1, 0), "Participant Printed Name: ___________________________")
    format_cell_text(t_c.cell(1, 1), "Official Position / Designation: ___________________________")
    format_cell_text(t_c.cell(2, 0), "Participant Signature: ___________________________\nDate: ___________________________")
    format_cell_text(t_c.cell(2, 1), "Researcher Signature: ___________________________\nDate: ___________________________")
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

def build_tool_03_content(doc):
    add_header(doc, "Stakeholder Needs Assessment Interview Guide",
               "Step 3: Semi-Structured Qualitative Protocol for Municipal & Provincial Personnel",
               "Step 3")
    add_callout(doc, "RESEARCH TIMELINE STAGE: STEP 3 (NEEDS ASSESSMENT & PROBLEM DISCOVERY)", [
        "When to Use: Conducted with municipal and provincial officers during requirements analysis to document workflow bottlenecks, disconnected records, and system requirements.",
        "Target Offices: Municipal Assessor's Office, Municipal Treasurer's Office, Provincial Assessor's Office (Pili).",
        "Student Researchers: Rizalina A. Ortega | John Kenneth P. Calubad | Christian Louise O. Asido",
        "Data Privacy Notice (RA 10173): All responses are confidential and used solely for academic research."
    ])

    t_det = doc.add_table(rows=3, cols=2)
    t_det.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_det)
    format_cell_text(t_det.cell(0, 0), "Date of Interview: ____________________", bold=True)
    format_cell_text(t_det.cell(0, 1), "Interviewer: ____________________", bold=True)
    format_cell_text(t_det.cell(1, 0), "Target Office: [  ] Assessor's Office    [  ] Treasurer's Office    [  ] Provincial Assessor")
    format_cell_text(t_det.cell(1, 1), "Recorded: [  ] Yes (Consented)    [  ] Notes Only")
    format_cell_text(t_det.cell(2, 0), "Designation / Position: ____________________")
    format_cell_text(t_det.cell(2, 1), "Years of Service in Office: ______ years")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    interview_parts = [
        ("Part A: Office Setup and Personnel Profile (All Participating Offices)", [
            "1. How many staff members currently work in your office, and what are their primary functional roles?",
            "2. On average, how many property tax transactions or client requests does your office handle on a typical day versus during peak deadline months?",
            "3. Which specific months or days of the week experience the highest volume of client visits, and what drives this surge?",
            "4. What are the most frequent counter requests, certifications, or official documents requested by property owners and other offices?",
        ]),
        ("Part B: Municipal Assessor's Office Workflows and Valuation Procedures", [
            "5. Could you walk us through the step-by-step process of handling a Transfer of Assessment, from initial document submission at the counter to final release of the new Tax Declaration?",
            "6. What specific documentary prerequisites must property owners submit for each major transaction?\n   a. New Tax Declaration (newly declared land, building/house, or machinery)\n   b. Transfer of Ownership (Deed of Sale, BIR CAR, transfer tax receipt, previous clearances)\n   c. Reassessment / Correction of Value (physical appraisal, reclassification)\n   d. Issuance of Clearances (Certificate of Land Holding, Certificate of No Property, Certificate of No Improvement)",
            "7. Once documents are submitted, how do personnel verify authenticity, and who conducts intermediate review and final approval prior to document release?",
            "8. How does your office compute assessed values and tax liabilities from start to finish? Is this done manually using physical appraisal tables, automated spreadsheets, or electronic software?",
            "9. Where and in what physical/electronic format are property records and assessment rolls currently stored? How do personnel retrieve an existing parcel record when an owner inquires?",
            "10. If your office currently utilizes a software system, what specific functions does it perform effectively, where does it fall short, and which procedural steps still strictly require paper handoffs?",
            "11. How frequently do clients experience repeated visits (pabalik-balik) due to incomplete paperwork or clerical errors, and which transaction types cause this most frequently?",
            "12. How does your office organize, protect, and retrieve decades-old legacy property archives (dating back to 1949 or earlier) against wear, deterioration, or loss?",
        ]),
        ("Part C: Municipal Treasurer's Office Workflows and Record Reconciliation", [
            "13. How does your office receive updated assessment rolls and tax declarations from the Assessor's Office? How frequently does this data exchange occur, and in what format?",
            "14. What is your step-by-step process for verifying whether a property tax has been paid before issuing official receipts or clearances? How do you link payment receipts to property records?",
            "15. On average, how many staff hours per week or month are spent manually cross-checking real property payment records against the municipal assessment roll?",
            "16. How does your office identify, track, and follow up delinquent or uncollected real property accounts?",
            "17. How are statutory prompt-payment discounts, interest penalties, and tax clearances computed and verified across multi-year assessment histories?",
            "18. Does your office utilize an electronic collection database or spreadsheet? Is it digitally synchronized with the Assessor's property records, or does it operate independently?",
            "19. What discrepancies or data mismatches happen most frequently between assessment records and payment rolls, and what protocol is followed to reconcile them?",
        ]),
        ("Part D: Provincial Assessor's Office Coordination and Approval (Pili Endorsement)", [
            "20. Which specific assessment transactions strictly require review, endorsement, and final statutory approval from the Provincial Assessor's Office?",
            "21. How are assessment packets physically transported from the municipality to the Provincial Capitol in Pili (physical courier, personnel delivery, mail), and what transit schedule is maintained?",
            "22. On average, how long does the provincial review and approval cycle take from the date of submission to physical return to the municipality, and what factors cause delays?",
            "23. What are the most common deficiencies or clerical errors that result in submitted packets being rejected or returned to the municipality for corrections?",
            "24. Is digital review or electronic endorsement of assessment documents permissible or viable under existing administrative standards? What technical and security safeguards would be required?",
            "25. How does your office currently communicate document processing status, endorsements, or required revisions back to municipal personnel?",
        ]),
        ("Part E: Inter-Office Coordination and Information Flow (Assessor, Treasurer, Provincial)", [
            "26. Which assessment and collection transactions demand close inter-office coordination, and through what communication channels is this handled (phone calls, personal office visits, paper memos, chat)?",
            "27. Do frontline staff currently possess visibility into where a document is queued across departments? How do personnel respond when a taxpayer requests a real-time status update?",
            "28. What procedural phase represents the single largest bottleneck or source of processing latency in the end-to-end workflow?",
            "29. What core capabilities should an integrated inter-office web portal provide to eliminate manual friction between the Assessor, Treasurer, and Provincial offices?",
        ]),
        ("Part F: Time-Motion Estimates (Self-Reported Perceptions)", [
            "30. For one typical Transfer of Assessment, estimate the average duration spent on each phase:\n   a. Counter requirement verification: ______ minutes\n   b. Archival record search and retrieval: ______ minutes\n   c. Valuation computation and encoding: ______ minutes\n   d. Supervisory review and municipal approval: ______ minutes\n   e. Inter-office transit and provincial approval queue: ______ days / hours\n   f. Total turnaround time from filing to release: ______ days",
            "31. For cross-checking tax payment rolls against property assessment ledgers, estimate the time spent per property record: ______ minutes.",
        ]),
        ("Part G: Stakeholder Feedback on the Proposed Technical Intervention", [
            "32. The proposed system provides a web-based platform that automates document routing, monitors transaction status in real-time, and synchronizes assessment and tax payment data. Which of your routine tasks would this simplify?",
            "33. What operational or technical concerns do you foresee regarding implementation (internet stability, staff digital literacy, data privacy, workflow adaptation)?",
            "34. What specific features or safeguards are indispensable for your office to adopt the proposed system?",
            "35. Would your office be open to participating in a simulated trial or parallel run alongside current office procedures using masked dummy data?",
            "36. Are there specific local ordinances, tax codes, or institutional procedures that the software engineering team must adhere to in designing the business logic?",
            "37. Do you have any additional recommendations or insights to guide the development of the system?",
        ]),
    ]

    for title, q_list in interview_parts:
        add_banner(doc, title)
        for q in q_list:
            p_q = doc.add_paragraph()
            p_q.paragraph_format.space_before = Pt(1.5)
            p_q.paragraph_format.space_after = Pt(2.5)
            r_q = p_q.add_run(q)
            r_q.font.size = Pt(8)
            r_q.font.color.rgb = COLOR_DARK
            r_q.font.name = "Arial"
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    add_callout(doc, "ACKNOWLEDGMENT & PARTICIPANT SIGN-OFF", [
        "I hereby confirm that I have participated in this stakeholder interview and provided feedback for academic research purposes under RA 10173.",
        "Participant Signature: ____________________________________        Date: ________________________",
        "Printed Name (Optional): __________________________________        Designation: ____________________"
    ])

def build_tool_04_content(doc):
    add_header(doc, "Time-Motion and Workflow Latency Log Sheet",
               "Step 4: Objective Timing Instrument for Baseline (Pre) vs. System (Post) Processing",
               "Step 4")
    add_callout(doc, "RESEARCH TIMELINE STAGE: STEP 4 (TIMING OBSERVATION - BASELINE & POST-TEST)", [
        "When to Use: Used in two phases: (1) Baseline Run before system deployment, and (2) Post-Test Run using the web synchronization system.",
        "Primary Quantitative Variables: (1) Document Processing Latency (days/hours) and (2) Record Reconciliation Time (minutes per property).",
        "Instructions: Log timestamps for each processing phase per case code (e.g., TOA-001). Do not record personal taxpayer details."
    ])

    t_g = doc.add_table(rows=3, cols=2)
    t_g.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_g)
    format_cell_text(t_g.cell(0, 0), "Case Transaction Code: ____________________", bold=True)
    format_cell_text(t_g.cell(0, 1), "Run Condition: [  ] BASELINE (Manual)    [  ] SYSTEM (Proposed)", bold=True)
    format_cell_text(t_g.cell(1, 0), "Transaction Type:\n[  ] Transfer of Assessment\n[  ] New Tax Declaration / Building\n[  ] Reassessment / Value Revision\n[  ] Tax Clearance / Certification\n[  ] Other: ____________________")
    format_cell_text(t_g.cell(1, 1), "Observed Department(s):\n[  ] Municipal Assessor's Office\n[  ] Municipal Treasurer's Office\n[  ] Provincial Assessor's Office (Pili)\n\nData Classification:\n[  ] Masked Dummy Data    [  ] Actual Workflow Log")
    format_cell_text(t_g.cell(2, 0), "Date of Observation: ____________________")
    format_cell_text(t_g.cell(2, 1), "Lead Observer / Researcher: ____________________")
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    add_banner(doc, "Table A: Phase-by-Phase Transaction Processing Latency")
    headers_a = ["Ph", "Operational Process Phase", "Dept", "Start Timestamp\n(Date / HH:MM)", "End Timestamp\n(Date / HH:MM)", "Duration\n(Min / Days)", "Bottleneck / Error Remarks"]
    steps_a = [
        ("1", "Counter intake & initial documentary completeness verification", "MAO"),
        ("2", "Property record search, manual ledger/eRPT parcel retrieval", "MAO"),
        ("3", "Field appraisal calculation, valuation encoding, FAAS prep", "MAO"),
        ("4", "Municipal Assessor review, endorsement & signature", "MAO"),
        ("5", "Inter-office document transfer to Municipal Treasurer queue", "Transit"),
        ("6", "Treasurer tax payment cross-check & clearance validation", "MTO"),
        ("7", "Physical/digital batch transport to Provincial Assessor (Pili)", "Transit"),
        ("8", "Provincial appraisal evaluation, title review & formal approval", "PAO"),
        ("9", "Return transit, final municipal encoding & Tax Declaration release", "MAO"),
    ]
    t_ta = doc.add_table(rows=len(steps_a) + 1, cols=7)
    t_ta.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_ta)
    
    for c_i, h in enumerate(headers_a):
        cell = t_ta.cell(0, c_i)
        set_cell_background(cell, HEX_MAROON)
        format_cell_text(cell, h, bold=True, size=7.5, color=RGBColor(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)

    for r_i, (ph, desc, dept) in enumerate(steps_a, start=1):
        format_cell_text(t_ta.cell(r_i, 0), ph, size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell_text(t_ta.cell(r_i, 1), desc, size=7.5)
        format_cell_text(t_ta.cell(r_i, 2), dept, size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        for c in range(3, 7):
            format_cell_text(t_ta.cell(r_i, c), "", size=7.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # Summary box
    t_sb = doc.add_table(rows=1, cols=3)
    t_sb.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_sb, color=HEX_MAROON_BORDER)
    format_cell_text(t_sb.cell(0, 0), "TOTAL ELAPSED TIME:\n______ Days ______ Hours ______ Min", bold=True, size=7.5)
    format_cell_text(t_sb.cell(0, 1), "NET ACTIVE WORK TIME:\n______ Minutes / Hours", bold=True, size=7.5)
    format_cell_text(t_sb.cell(0, 2), "QUEUE / TRANSIT LATENCY:\n______ Days / Hours", bold=True, size=7.5)
    for c in t_sb.rows[0].cells:
        set_cell_background(c, HEX_LIGHT_RED)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    add_banner(doc, "Table B: Record Reconciliation and Payment Verification Metrics")
    recon_rows = [
        ("Total Property Records Checked in Batch", "______ records", "Sample batch of property accounts cross-referenced."),
        ("Batch Verification Start Timestamp", "Date: ________ Time: ____", "Clock time upon initiation of record cross-checking."),
        ("Batch Verification Completion Timestamp", "Date: ________ Time: ____", "Clock time upon resolution of the final record in batch."),
        ("Total Batch Execution Duration", "______ minutes", "Net active minutes elapsed to reconcile complete batch."),
        ("Average Reconciliation Latency per Record", "______ min/record", "Average active minutes elapsed per individual record."),
        ("Discrepancies / Data Mismatches Detected", "______ mismatches", "Frequency of PIN, owner, or value mismatches identified."),
        ("Reconciliation Mechanism Employed", "[  ] Manual Ledgers / Excel\n[  ] Proposed Sync Engine", "Comparative verification method utilized during observation.")
    ]
    t_tb = doc.add_table(rows=len(recon_rows) + 1, cols=3)
    t_tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_tb)
    
    headers_b = ["Batch Reconciliation Parameter", "Observed Metric Value", "Operational Context & Methodological Notes"]
    for c_i, h in enumerate(headers_b):
        cell = t_tb.cell(0, c_i)
        set_cell_background(cell, HEX_MAROON)
        format_cell_text(cell, h, bold=True, size=7.5, color=RGBColor(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)

    for r_i, (param, val, notes) in enumerate(recon_rows, start=1):
        format_cell_text(t_tb.cell(r_i, 0), param, bold=True, size=7.5)
        format_cell_text(t_tb.cell(r_i, 1), val, size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell_text(t_tb.cell(r_i, 2), notes, size=7.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    add_banner(doc, "Table C: Comparative Evaluation Summary and Reduction Analysis")
    comp_rows = [
        ("Document Processing Latency (Transfer of Assessment, days)", "______ days", "______ days", "______ days", "______ %"),
        ("Building Assessment End-to-End Latency (days)", "______ days", "______ days", "______ days", "______ %"),
        ("Record Reconciliation Latency per Property (minutes)", "______ min", "______ min", "______ min", "______ %"),
        ("Frequency of In-Person Status Inquiries per Transaction", "______ visits", "______ visits", "______ visits", "______ %"),
    ]
    t_tc = doc.add_table(rows=len(comp_rows) + 1, cols=5)
    t_tc.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_tc)
    
    headers_c = ["Performance Indicator Metric", "Baseline Mean (T0)", "System Mean (T1)", "Absolute Difference (Δ)", "Percentage Reduction (% Red)"]
    for c_i, h in enumerate(headers_c):
        cell = t_tc.cell(0, c_i)
        set_cell_background(cell, HEX_MAROON)
        format_cell_text(cell, h, bold=True, size=7.5, color=RGBColor(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)

    for r_i, (metric, b_val, s_val, diff, red) in enumerate(comp_rows, start=1):
        format_cell_text(t_tc.cell(r_i, 0), metric, size=7.5)
        format_cell_text(t_tc.cell(r_i, 1), b_val, size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell_text(t_tc.cell(r_i, 2), s_val, size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell_text(t_tc.cell(r_i, 3), diff, size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell_text(t_tc.cell(r_i, 4), red, size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    add_callout(doc, "OBSERVER VERIFICATION & SUBMISSION", [
        "I hereby certify that the timestamps and duration metrics recorded in this observation log accurately reflect the observed workflow times.",
        "Observer Signature: ____________________________________        Verification Date: ________________________",
        "Printed Name: __________________________________________        Office / Station: ____________________"
    ])

def build_tool_05_content(doc):
    add_header(doc, "Software Quality Evaluation Survey (ISO/IEC 25010)",
               "Step 5: Post-Implementation System Quality Rating Instrument",
               "Step 5")
    add_callout(doc, "RESEARCH TIMELINE STAGE: STEP 5 (POST-DEVELOPMENT SYSTEM EVALUATION)", [
        "When to Use: Administered after municipal staff and IT experts test the developed web portal and document routing engine.",
        "Target Respondent Groups: Municipal Assessor staff, Municipal Treasurer staff, Provincial Assessor officers, and IT Experts.",
        "Standard 5-Point Likert Scale: [ 5 ] Excellent  |  [ 4 ] Very Good  |  [ 3 ] Good  |  [ 2 ] Fair  |  [ 1 ] Poor",
        "Directions: Following hands-on interaction with the system, please mark the rating (1 to 5) corresponding to each evaluation statement."
    ])

    t_p = doc.add_table(rows=2, cols=2)
    t_p.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_p)
    format_cell_text(t_p.cell(0, 0), "Respondent Name (Optional): ___________________________", bold=True)
    format_cell_text(t_p.cell(0, 1), "Date of Evaluation: ___________________________", bold=True)
    format_cell_text(t_p.cell(1, 0), "Institutional Department / Group:\n[  ] Municipal Assessor's Office\n[  ] Municipal Treasurer's Office\n[  ] Provincial Assessor's Office (Pili)\n[  ] IT Expert / Systems Evaluator")
    format_cell_text(t_p.cell(1, 1), "Designation / Role: ___________________________\n\nYears of Service: [  ] < 1 yr  [  ] 1-5 yrs  [  ] 6-10 yrs  [  ] > 10 yrs\n\nComputer Literacy: [  ] Basic  [  ] Intermediate  [  ] Advanced")
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    iso_sections = [
        ("A. Functional Suitability", [
            ("1", "The system allows personnel to generate, update, and manage property assessment records according to statutory valuation rules."),
            ("2", "The automated document routing engine dispatches assessment packets accurately across Assessor, Treasurer, and Provincial queues."),
            ("3", "The payment synchronization module matches municipal treasury payment entries with corresponding property assessment records reliably."),
            ("4", "The platform detects and flags discrepancies or duplicate entries between assessment rolls and collection ledgers."),
            ("5", "Automated computations for assessed values, basic tax liabilities, and Special Education Fund (SEF) dues are precise and error-free."),
            ("6", "The functional scope covers the core requirements necessary to eliminate manual paper transfers and physical reconciliation."),
        ]),
        ("B. Performance Efficiency", [
            ("7", "System interfaces, data tables, and interactive dashboards load promptly without noticeable latency."),
            ("8", "Document routing status transitions update in real time across departmental monitors upon completion of procedural actions."),
            ("9", "Property index search queries (PIN, owner name, ARP number) return matching parcel records within three seconds."),
            ("10", "End-to-end transaction duration (e.g., Transfer of Assessment) is noticeably shorter compared to legacy manual workflows."),
            ("11", "System responsiveness remains consistent and stable when multiple municipal staff access the platform concurrently."),
        ]),
        ("C. Usability", [
            ("12", "User interface layouts and visual hierarchies are clear, structured, and straightforward to navigate."),
            ("13", "Frontline staff can comprehend and master primary transactional workflows with minimal technical training."),
            ("14", "Routine operational tasks (record intake, validation, clearance issuance) can be accomplished within a few user actions."),
            ("15", "The real-time document tracking dashboard transparently displays current document location, queue status, and assigned personnel."),
            ("16", "System feedback, input validations, and error notifications provide actionable and understandable instructions to the user."),
            ("17", "The software substantially simplifies inter-office coordination and eliminates tedious manual verification handoffs."),
        ]),
        ("D. Security", [
            ("18", "The system enforces secure user authentication protocols and prevents unauthorized system entry."),
            ("19", "Role-Based Access Control (RBAC) restricts data views and transactional privileges strictly to authorized municipal roles."),
            ("20", "Automated session timeouts terminate inactive user sessions to prevent unauthorized terminal access."),
            ("21", "An unalterable audit trail accurately logs all document modifications, status updates, and supervisory approvals with timestamps."),
            ("22", "Citizen property records and assessment rolls are effectively protected in accordance with the Data Privacy Act (RA 10173)."),
        ]),
        ("E. Reliability", [
            ("23", "The web application functions stably without sudden runtime crashes, deadlocks, or unhandled system exceptions."),
            ("24", "Recorded property assessments, payment logs, and file attachments remain intact and uncorrupted following logout and re-authentication."),
            ("25", "Atomic database transactions guarantee that interrupted or failed actions do not result in corrupted, partial, or orphaned records."),
            ("26", "The system maintains continuous availability during standard municipal operating hours."),
            ("27", "The platform recovers gracefully from intermittent network disconnections without causing transactional data loss."),
        ]),
        ("F. Compatibility (Platform Interoperability)", [
            ("28", "The web application functions identically and consistently across major modern web browsers (Chrome, Edge, Firefox)."),
            ("29", "The system operates smoothly on standard municipal office desktop computers without requiring proprietary hardware upgrades."),
        ]),
    ]

    for sec_title, items in iso_sections:
        add_banner(doc, sec_title)
        t_s = doc.add_table(rows=len(items) + 1, cols=7)
        t_s.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t_s)
        
        headers_s = ["#", "Evaluation Criterion Statement", "5 (Exc)", "4 (V.Gd)", "3 (Good)", "2 (Fair)", "1 (Poor)"]
        for c_i, h in enumerate(headers_s):
            cell = t_s.cell(0, c_i)
            set_cell_background(cell, HEX_MAROON)
            format_cell_text(cell, h, bold=True, size=7.5, color=RGBColor(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)

        for r_i, (num, stmt) in enumerate(items, start=1):
            format_cell_text(t_s.cell(r_i, 0), num, size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)
            format_cell_text(t_s.cell(r_i, 1), stmt, size=7.5)
            for c in range(2, 7):
                format_cell_text(t_s.cell(r_i, c), "[   ]", size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    add_callout(doc, "QUALITATIVE FEEDBACK & USER ACCEPTANCE", [
        "1. Which feature or capability provided the greatest operational benefit to your workflow?\n____________________________________________________________________________________________________________",
        "2. What operational difficulties or technical bugs did you encounter during testing?\n____________________________________________________________________________________________________________",
        "3. What modifications or additional features would you recommend for future versions?\n____________________________________________________________________________________________________________",
        "4. Would you endorse adopting the proposed system in your office?   [  ] Yes    [  ] No    [  ] Conditional\nReason: ___________________________________________________________________________________________________"
    ])

    add_callout(doc, "EVALUATOR VERIFICATION & SUBMISSION", [
        "I hereby certify that I have thoroughly evaluated the software features and that the ratings provided above reflect my genuine operational and technical evaluation of the platform.",
        "Evaluator Signature: ____________________________________        Submission Date: ________________________",
        "Printed Name: __________________________________________        Official Designation: ____________________"
    ])

def generate_all():
    print("Generating Combined Master Document (All Tools Step 1 to 5)...")
    doc_combined = Document()
    
    # Set page margins to 0.5 in
    sections = doc_combined.sections
    for s in sections:
        s.top_margin = Inches(0.5)
        s.bottom_margin = Inches(0.5)
        s.left_margin = Inches(0.5)
        s.right_margin = Inches(0.5)

    # Tool 1
    build_tool_01_content(doc_combined)
    doc_combined.add_page_break()

    # Tool 2
    build_tool_02_content(doc_combined)
    doc_combined.add_page_break()

    # Tool 3
    build_tool_03_content(doc_combined)
    doc_combined.add_page_break()

    # Tool 4
    build_tool_04_content(doc_combined)
    doc_combined.add_page_break()

    # Tool 5
    build_tool_05_content(doc_combined)

    combined_path = os.path.join(DOCX_DIR, "Research_Instruments_All_Tools_Step1_to_Step5.docx")
    doc_combined.save(combined_path)
    print("Saved combined document:", combined_path)

    # Also generate the 5 individual files in Individual_Files folder
    tools = [
        ("01_Expert_Validation_Sheet_Pre_Deployment.docx", build_tool_01_content),
        ("02_Participant_Consent_and_Data_Privacy_Agreement.docx", build_tool_02_content),
        ("03_Stakeholder_Needs_Assessment_Interview_Guide.docx", build_tool_03_content),
        ("04_Time_Motion_and_Workflow_Latency_Log_Sheet.docx", build_tool_04_content),
        ("05_ISO_IEC_25010_Software_Quality_Evaluation_Survey.docx", build_tool_05_content),
    ]

    for fname, builder_func in tools:
        doc_indiv = Document()
        for s in doc_indiv.sections:
            s.top_margin = Inches(0.5)
            s.bottom_margin = Inches(0.5)
            s.left_margin = Inches(0.5)
            s.right_margin = Inches(0.5)
        builder_func(doc_indiv)
        indiv_path = os.path.join(INDIV_DIR, fname)
        doc_indiv.save(indiv_path)
        print("Saved individual document:", indiv_path)

if __name__ == "__main__":
    generate_all()
    print("ALL DOCX FILES SUCCESSFULLY GENERATED!")
